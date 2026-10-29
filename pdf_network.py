#!/usr/bin/env python3
import datetime
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import pymupdf
from PySide6.QtCore import Qt, QUrl, QEvent
from PySide6.QtGui import QIcon, QDesktopServices
from PySide6.QtWidgets import (
    QApplication, QFileDialog, QHBoxLayout, QInputDialog, QLabel,
    QLineEdit, QListWidget, QListWidgetItem, QMainWindow, QMenu,
    QMessageBox, QPushButton, QSplitter, QStackedWidget, QTextBrowser,
    QTreeWidgetItem, QVBoxLayout, QWidget, QComboBox
)

from database import Database
from widgets.bookmark_tree import BookmarkTree
from widgets.pdf_viewer import PDFViewer
from lang import t, set_lang

sys.dont_write_bytecode = True

PROJE_ANA_DIZINI = Path(os.path.realpath(__file__)).parent
os.chdir(PROJE_ANA_DIZINI)

class NoteBrowser(QTextBrowser):
    """Not görüntüleyici ve çift tıkla düzenleme desteği."""
    def __init__(self, parent=None):
        super().__init__(parent)
        self.main_window = None

    def mouseDoubleClickEvent(self, event):
        if self.main_window:
            self.main_window.enable_note_editing()
        super().mouseDoubleClickEvent(event)

class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.db = Database()

        icon_path = PROJE_ANA_DIZINI / "icon.svg"
        if icon_path.exists():
            self.setWindowIcon(QIcon(str(icon_path)))

        self.viewer = PDFViewer()
        self.tree = BookmarkTree()

        # UI Elemanları
        self.page_label = QLabel()
        self.zoom_label = QLabel()
        self.btn_open = QPushButton()
        self.btn_group = QPushButton()
        self.btn_prev = QPushButton()
        self.btn_next = QPushButton()
        self.btn_zoom_in = QPushButton()
        self.btn_zoom_out = QPushButton()
        self.btn_toggle = QPushButton()
        self.btn_print = QPushButton()
        self.btn_export = QPushButton()
        self.note_label = QLabel()
        self.btn_save_note = QPushButton()
        self.btn_search = QPushButton()
        self.dash_title = QLabel()
        self.dash_recent_lbl = QLabel()

        # Sinyaller
        self.viewer.zoom_changed.connect(self.update_zoom_label)
        self.viewer.page_changed.connect(self.update_page_label)
        self.viewer.bookmarkRequested.connect(self.add_bookmark_from_selection)
        self.viewer.highlightRequested.connect(self.create_highlight)
        self.viewer.deleteHighlightRequested.connect(self.delete_highlight)
        self.viewer.exportTableRequested.connect(self.export_table_to_excel)

        self.settings_file = PROJE_ANA_DIZINI / "data" / "settings.json"
        self.selected_group_id = None
        self.active_bookmark_id = None
        self.pdf_history = []
        self.current_pdf = None

        self.setup_ui()

        self.load_settings()
        self.metinleri_guncelle()
        self.load_groups()

        self.tree.setContextMenuPolicy(Qt.CustomContextMenu)
        self.tree.customContextMenuRequested.connect(self.show_tree_menu)
        self.tree.itemClicked.connect(self.on_tree_item_clicked)

    def metinleri_guncelle(self):
        base_title = t("app_title")
        if self.current_pdf:
            self.setWindowTitle(f"{base_title} - {Path(self.current_pdf).name}")
        else:
            self.setWindowTitle(base_title)

        self.update_page_label()
        self.update_zoom_label()

        self.page_edit.setPlaceholderText(t("ph_page"))
        self.btn_open.setText(t("btn_open"))
        self.btn_group.setText(t("btn_group"))
        self.btn_prev.setText(t("btn_prev"))
        self.btn_next.setText(t("btn_next"))
        self.btn_zoom_in.setText(t("btn_zoom_in"))
        self.btn_zoom_out.setText(t("btn_zoom_out"))
        self.btn_toggle.setText(t("btn_toggle"))
        self.btn_print.setText(t("btn_print"))
        self.btn_export.setText(t("btn_export"))

        self.search_box.setPlaceholderText(t("ph_search_tree"))
        self.note_label.setText(t("lbl_note"))
        self.note_edit.setPlaceholderText(t("ph_note"))
        self.btn_save_note.setText(t("btn_save_note"))

        self.pdf_search.setPlaceholderText(t("ph_search_pdf"))
        self.btn_search.setText(t("btn_search_find"))

        self.dash_title.setText(t("dash_title"))
        self.dash_recent_lbl.setText(t("dash_recent_title"))
        self.tree.setHeaderLabel(t("tree_header"))

        self.update_dashboard_ui(self.pdf_history)
        self.load_groups()

    def lang_degistir(self, index):
        if index == 0:
            set_lang("tr")
        else:
            set_lang("en")
        self.metinleri_guncelle()
        self.save_settings()

    def setup_ui(self):
        central = QWidget()
        root = QVBoxLayout()
        toolbar = QHBoxLayout()

        self.page_edit = QLineEdit()
        self.page_edit.setFixedWidth(80)
        self.page_edit.returnPressed.connect(self.goto_page)

        toolbar.setContentsMargins(0, 0, 0, 0)
        toolbar.setSpacing(4)

        self.btn_open.clicked.connect(self.open_pdf)
        self.btn_prev.clicked.connect(self.previous_page)
        self.btn_next.clicked.connect(self.next_page)
        self.btn_zoom_in.clicked.connect(self.zoom_in)
        self.btn_zoom_out.clicked.connect(self.zoom_out)
        self.btn_group.clicked.connect(self.add_group)
        self.btn_toggle.clicked.connect(self.toggle_bookmarks)
        self.btn_print.clicked.connect(self.print_pdf)
        self.btn_export.clicked.connect(self.export_bookmarks_to_markdown)

        self.combo_lang = QComboBox()
        self.combo_lang.addItems(["🇹🇷 TR", "🇬🇧 EN"])
        self.combo_lang.setMinimumWidth(85)
        self.combo_lang.currentIndexChanged.connect(self.lang_degistir)

        for btn in (self.btn_print, self.btn_open, self.btn_prev, self.btn_next,
                    self.btn_zoom_in, self.btn_zoom_out, self.btn_group,
                    self.btn_toggle, self.btn_export, self.page_label,
                    self.zoom_label, self.combo_lang):
            btn.setMaximumHeight(28)

        toolbar.addWidget(self.btn_open)
        toolbar.addWidget(self.btn_print)
        toolbar.addWidget(self.btn_export)
        toolbar.addWidget(self.btn_group)
        toolbar.addWidget(self.btn_prev)
        toolbar.addWidget(self.btn_next)
        toolbar.addWidget(self.page_edit)
        toolbar.addWidget(self.btn_zoom_in)
        toolbar.addWidget(self.btn_zoom_out)
        toolbar.addWidget(self.btn_toggle)
        toolbar.addStretch()
        toolbar.addWidget(self.zoom_label)
        toolbar.addWidget(self.page_label)
        toolbar.addWidget(self.combo_lang)

        # Sol panel
        self.search_box = QLineEdit()
        self.search_box.textChanged.connect(self.filter_tree)

        self.left_panel = QWidget()
        self.left_panel.setMinimumWidth(200)
        left_layout = QVBoxLayout(self.left_panel)
        left_layout.setContentsMargins(0, 0, 0, 0)
        left_layout.setSpacing(4)

        self.note_label.setStyleSheet("font-weight: bold; margin-top: 5px;")

        self.note_edit = NoteBrowser()
        self.note_edit.main_window = self
        self.note_edit.setReadOnly(True)
        self.note_edit.setOpenLinks(False)
        self.note_edit.setMaximumHeight(150)
        self.note_edit.setEnabled(False)
        self.note_edit.anchorClicked.connect(self.handle_knowledge_link)

        self.btn_save_note.setMaximumHeight(28)
        self.btn_save_note.setEnabled(False)
        self.btn_save_note.clicked.connect(self.save_current_bookmark_note)

        left_layout.addWidget(self.search_box)
        left_layout.addWidget(self.tree, 1)
        left_layout.addWidget(self.note_label)
        left_layout.addWidget(self.note_edit)
        left_layout.addWidget(self.btn_save_note)

        # Arama paneli
        self.pdf_search = QLineEdit()
        self.pdf_search.returnPressed.connect(self.search_in_pdf)
        self.btn_search.clicked.connect(self.search_in_pdf)

        self.search_results = QListWidget()
        self.search_results.itemDoubleClicked.connect(self.goto_search_result)

        search_panel = QWidget()
        search_layout = QVBoxLayout(search_panel)
        search_layout.setContentsMargins(0, 0, 0, 0)
        search_layout.addWidget(self.pdf_search)
        search_layout.addWidget(self.btn_search)
        search_layout.addWidget(self.search_results)

        # Splitter & Stack
        self.splitter = QSplitter()
        self.splitter.addWidget(self.left_panel)

        self.right_stack = QStackedWidget()

        self.dashboard_widget = QWidget()
        dash_layout = QVBoxLayout(self.dashboard_widget)
        dash_layout.setContentsMargins(20, 20, 20, 20)
        dash_layout.setSpacing(15)

        self.dash_title.setStyleSheet("font-size: 20px; font-weight: bold; color: #2c3e50;")
        dash_layout.addWidget(self.dash_title)

        self.stats_label = QLabel()
        self.stats_label.setStyleSheet("font-size: 13px; color: #7f8c8d; font-style: italic;")
        dash_layout.addWidget(self.stats_label)

        dash_layout.addWidget(self.dash_recent_lbl)

        self.history_list = QListWidget()
        self.history_list.setAlternatingRowColors(True)
        self.history_list.setStyleSheet("""
            QListWidget {
                font-size: 14px;
                padding: 5px;
                background-color: #ffffff;
            }
            QListWidget::item {
                padding: 8px;
                border-bottom: 1px solid #e0e0e0;
            }
            QListWidget::item:alternate {
                background-color: #f4f6f9;
            }
            QListWidget::item:selected {
                background-color: #d0e6f5;
                color: black;
            }
        """)
        self.history_list.itemDoubleClicked.connect(self.open_from_dashboard)
        dash_layout.addWidget(self.history_list, 1)

        self.right_stack.addWidget(self.dashboard_widget)
        self.right_stack.addWidget(self.viewer)

        self.splitter.addWidget(self.right_stack)
        self.splitter.addWidget(search_panel)

        self.splitter.setSizes([250, 900, 300])
        self.splitter.setStretchFactor(0, 0)
        self.splitter.setStretchFactor(1, 1)

        root.addLayout(toolbar)
        root.addWidget(self.splitter, 1)

        central.setLayout(root)
        self.setCentralWidget(central)
        self.resize(1400, 900)

    def on_tree_item_clicked(self, item):
        group_id = item.data(0, Qt.UserRole)
        if group_id is not None:
            self.selected_group_id = group_id
        self.open_bookmark(item)

    def open_pdf(self):
        default_dir = str(Path.home())
        file_filter = t("filter_supported")
        filename, _ = QFileDialog.getOpenFileName(self, t("btn_open"), default_dir, file_filter)

        if not filename:
            return
        self.current_pdf = filename
        self.viewer.open_pdf(filename)
        self.update_page_label()
        self.load_pdf_bookmarks()
        self.setWindowTitle(f"{t('app_title')} - {Path(filename).name}")
        self.right_stack.setCurrentIndex(1)

    def load_pdf_bookmarks(self):
        pdf_root_name = t("pdf_bookmarks_root")
        for i in range(self.tree.topLevelItemCount()):
            item = self.tree.topLevelItem(i)
            if item.text(0) in (pdf_root_name, "📘 Belge Yer İmleri", "📘 Document Bookmarks"):
                self.tree.takeTopLevelItem(i)
                break

        if not self.viewer.doc:
            return

        toc = self.viewer.doc.get_toc()
        self.pdf_root = QTreeWidgetItem([pdf_root_name])
        self.tree.addTopLevelItem(self.pdf_root)
        parents = {0: self.pdf_root}

        for level, title, page in toc:
            item = QTreeWidgetItem([title])
            item.setData(0, Qt.UserRole + 100, page - 1)
            parent_item = parents.get(level - 1, self.pdf_root)
            parent_item.addChild(item)
            parents[level] = item
        self.pdf_root.setExpanded(True)

    def zoom_in(self):
        self.viewer.zoom_in()
        self.update_zoom_label()

    def zoom_out(self):
        self.viewer.zoom_out()
        self.update_zoom_label()

    def update_zoom_label(self):
        value = int(self.viewer.zoom * 100)
        self.zoom_label.setText(t("lbl_zoom").format(value))

    def next_page(self):
        self.viewer.next_page()
        self.update_page_label()

    def previous_page(self):
        self.viewer.previous_page()
        self.update_page_label()

    def update_page_label(self):
        if not self.viewer.doc:
            self.page_label.setText(t("lbl_page_empty"))
            self.page_edit.clear()
            return

        current = self.viewer.current_page + 1
        total = len(self.viewer.doc)
        self.page_label.setText(t("lbl_page").format(current, total))
        self.page_edit.setText(str(current))

    def add_group(self):
        name, ok = QInputDialog.getText(self, t("dlg_new_group"), t("dlg_group_name"))
        if not ok or not name.strip():
            return
        self.db.add_group(name.strip())
        self.load_groups()

    def load_groups(self):
        self.tree.clear()
        groups = self.db.get_groups()
        nodes = {}

        for gid, name, parent_id in groups:
            item = QTreeWidgetItem([name])
            item.setData(0, Qt.UserRole, gid)
            nodes[gid] = item

        for gid, name, parent_id in groups:
            item = nodes[gid]
            if parent_id and parent_id in nodes:
                nodes[parent_id].addChild(item)
            else:
                self.tree.addTopLevelItem(item)

        bookmarks = self.db.get_bookmarks()
        for bookmark in bookmarks:
            bid, title, pdf_path, page, group_id = bookmark[0], bookmark[1], bookmark[2], bookmark[3], bookmark[4]
            if group_id in nodes:
                bookmark_item = QTreeWidgetItem([title])
                bookmark_item.setData(0, Qt.UserRole + 1, bid)
                nodes[group_id].addChild(bookmark_item)

        if self.viewer.doc:
            self.load_pdf_bookmarks()

    def toggle_bookmarks(self):
        sizes = self.splitter.sizes()
        if sizes[0] > 0:
            self._last_left_width = sizes[0]
            sizes[0] = 0
        else:
            sizes[0] = getattr(self, "_last_left_width", 250)
        self.splitter.setSizes(sizes)

    def save_settings(self):
        sizes = self.splitter.sizes()
        is_visible = sizes[0] > 0
        if not is_visible and hasattr(self, "_last_left_width"):
            sizes[0] = self._last_left_width

        history = getattr(self, "pdf_history", [])
        if self.current_pdf:
            history = [item for item in history if (item.get("path") if isinstance(item, dict) else item) != self.current_pdf]
            history.insert(0, {
                "path": self.current_pdf,
                "page": self.viewer.current_page,
                "date": datetime.datetime.now().strftime("%Y-%m-%d %H:%M")
            })
            history = history[:5]

        data = {
            "window_width": self.width(),
            "window_height": self.height(),
            "splitter_sizes": sizes,
            "bookmarks_visible": is_visible,
            "pdf_history": history,
            "last_active_index": self.right_stack.currentIndex(),
            "language": "tr" if self.combo_lang.currentIndex() == 0 else "en"
        }
        self.settings_file.parent.mkdir(parents=True, exist_ok=True)
        with open(self.settings_file, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=4)

    def load_settings(self):
        if not self.settings_file.exists():
            self.update_dashboard_ui([])
            return
        try:
            with open(self.settings_file, "r", encoding="utf-8") as f:
                data = json.load(f)

            kayitli_dil = data.get("language", "tr")
            self.combo_lang.blockSignals(True)
            if kayitli_dil == "en":
                self.combo_lang.setCurrentIndex(1)
                set_lang("en")
            else:
                self.combo_lang.setCurrentIndex(0)
                set_lang("tr")
            self.combo_lang.blockSignals(False)

            self.resize(data.get("window_width", 1400), data.get("window_height", 900))
            splitter_sizes = data.get("splitter_sizes", [250, 900, 300])
            is_visible = data.get("bookmarks_visible", True)

            if not is_visible:
                self._last_left_width = splitter_sizes[0]
                splitter_sizes[0] = 0
            else:
                self._last_left_width = splitter_sizes[0]

            self.splitter.setSizes(splitter_sizes)
            self.pdf_history = data.get("pdf_history", [])
            self.update_dashboard_ui(self.pdf_history)
            self.right_stack.setCurrentIndex(0)

            if self.pdf_history:
                last_item = self.pdf_history[0]
                pdf_path = last_item.get("path") if isinstance(last_item, dict) else last_item
                if pdf_path and Path(pdf_path).exists():
                    self.current_pdf = pdf_path

        except Exception as e:
            print("Settings error:", e)
            self.update_dashboard_ui([])

    def update_dashboard_ui(self, history):
        self.history_list.clear()
        try:
            total_bookmarks = self.db.get_bookmark_count()
            total_groups = self.db.get_group_count()
            self.stats_label.setText(t("dash_stats").format(total_bookmarks, total_groups))
        except Exception:
            self.stats_label.setText(t("dash_stats_error"))

        if not history:
            item = QListWidgetItem(t("dash_empty"))
            item.setFlags(Qt.NoItemFlags)
            self.history_list.addItem(item)
            return

        for data in history:
            if not isinstance(data, dict):
                continue
            path = data.get("path", "")
            page = data.get("page", 0) + 1
            date = data.get("date", "-")
            file_name = Path(path).name

            display_text = t("dash_item").format(file_name, path, page, date)
            list_item = QListWidgetItem(display_text)
            list_item.setData(Qt.UserRole, data)
            self.history_list.addItem(list_item)

    def open_from_dashboard(self, item):
        data = item.data(Qt.UserRole)
        if not data:
            return

        pdf_path = data.get("path")
        page = data.get("page", 0)

        if not pdf_path or not Path(pdf_path).exists():
            QMessageBox.warning(self, t("msg_error"), t("err_file_missing"))
            return

        self.current_pdf = pdf_path
        self.viewer.open_pdf(pdf_path)
        self.viewer.goto_page(page)
        self.update_page_label()
        self.load_pdf_bookmarks()
        self.right_stack.setCurrentIndex(1)
        self.setWindowTitle(f"{t('app_title')} - {Path(pdf_path).name}")

    def closeEvent(self, event):
        self.save_settings()
        super().closeEvent(event)

    def search_in_pdf(self):
        text = self.pdf_search.text().strip()
        if not text:
            return

        pages = self.viewer.search_text(text)
        grouped = {}

        for page, rect in pages:
            if page not in grouped:
                grouped[page] = {"count": 0, "rects": []}
            grouped[page]["count"] += 1
            grouped[page]["rects"].append((page, rect))

        self.search_results.clear()
        total_matches = len(pages)
        self.search_results.addItem(t("search_matches").format(total_matches))

        for page in sorted(grouped):
            count = grouped[page]["count"]
            rects = grouped[page]["rects"]
            item = QListWidgetItem(t("search_page_result").format(page + 1, count))
            item.setData(Qt.UserRole, (page, rects))
            self.search_results.addItem(item)

    def goto_search_result(self, item):
        if item.data(Qt.UserRole) is None:
            return
        page, rects = item.data(Qt.UserRole)
        if not isinstance(rects, list) or len(rects) == 0:
            return

        self.viewer.goto_page(page)
        self.viewer.set_highlight(rects)

        first_item = rects[0]
        if isinstance(first_item, tuple) and len(first_item) == 2:
            self.viewer.active_search_rect = first_item[1]
        else:
            self.viewer.active_search_rect = first_item

        self.viewer.render_page()
        self.update_page_label()

    def filter_tree(self, text):
        text = text.lower().strip()
        for i in range(self.tree.topLevelItemCount()):
            item = self.tree.topLevelItem(i)
            self.filter_item(item, text)

    def filter_item(self, item, text):
        visible = text in item.text(0).lower()
        child_visible = False
        for i in range(item.childCount()):
            if self.filter_item(item.child(i), text):
                child_visible = True
        show = visible or child_visible or text == ""
        item.setHidden(not show)
        if child_visible and text != "":
            item.setExpanded(True)
        return show

    def show_tree_menu(self, pos):
        item = self.tree.itemAt(pos)
        menu = QMenu()
        add_bookmark = menu.addAction(t("menu_add_bookmark"))
        rename_bookmark = menu.addAction(t("menu_rename_bookmark"))
        delete_bookmark = menu.addAction(t("menu_del_bookmark"))
        add_subgroup = menu.addAction(t("menu_add_subgroup"))
        rename_group = menu.addAction(t("menu_rename_group"))
        delete_group = menu.addAction(t("menu_del_group"))
        add_cross_link = menu.addAction(t("menu_add_link"))

        action = menu.exec(self.tree.mapToGlobal(pos))
        if action == add_bookmark:
            self.add_bookmark(item)
        elif action == rename_bookmark:
            self.rename_bookmark(item)
        elif action == delete_bookmark:
            self.delete_bookmark(item)
        elif action == add_subgroup:
            self.add_subgroup(item)
        elif action == rename_group:
            self.rename_group(item)
        elif action == delete_group:
            self.delete_group_recursive(item)
        elif action == add_cross_link:
            self.add_cross_pdf_link(item)

    def add_bookmark(self, item):
        if not self.current_pdf or not item:
            return
        current_page = self.viewer.current_page
        group_id = item.data(0, Qt.UserRole)
        title, ok = QInputDialog.getText(self, t("dlg_new_bookmark"), t("dlg_bookmark_name"))
        if not ok or not title.strip():
            return
        self.db.add_bookmark(title.strip(), self.current_pdf, current_page, group_id)
        self.load_groups()
        self.viewer.goto_page(current_page)
        self.update_page_label()

    def delete_bookmark(self, item):
        if not item: return
        bookmark_id = item.data(0, Qt.UserRole + 1)
        if not bookmark_id: return
        self.db.delete_bookmark(bookmark_id)
        self.load_groups()

    def rename_bookmark(self, item):
        if not item: return
        bookmark_id = item.data(0, Qt.UserRole + 1)
        if not bookmark_id: return
        name, ok = QInputDialog.getText(self, t("dlg_rename_bookmark"), t("dlg_new_name"), text=item.text(0))
        if ok and name.strip():
            self.db.rename_bookmark(bookmark_id, name.strip())
            self.load_groups()

    def open_bookmark(self, item):
        if QApplication.mouseButtons() == Qt.RightButton:
            return
        pdf_page = item.data(0, Qt.UserRole + 100)
        if pdf_page is not None:
            self.viewer.goto_page(pdf_page)
            self.update_page_label()
            self.right_stack.setCurrentIndex(1)
            return
        bookmark_id = item.data(0, Qt.UserRole + 1)
        if not bookmark_id:
            self.active_bookmark_id = None
            self.note_edit.clear()
            self.note_edit.setEnabled(False)
            self.btn_save_note.setEnabled(False)
            return

        bookmark = self.db.get_bookmark(bookmark_id)
        if not bookmark:
            return

        self.active_bookmark_id = bookmark_id
        pdf_path = bookmark[2]
        page = bookmark[3]
        x = bookmark[5] if len(bookmark) > 5 else 0
        y = bookmark[6] if len(bookmark) > 6 else 0
        current_note = bookmark[7] if len(bookmark) > 7 and bookmark[7] else ""

        if "[[" in current_note and "]]" in current_note and "#" in current_note:
            try:
                link_match = re.search(r'\[\[([^\]]+)\]\]', current_note)
                if link_match:
                    _, page_str = link_match.group(1).split("#", 1)
                    clean_page_str = "".join(filter(str.isdigit, page_str))
                    if clean_page_str:
                        page = int(clean_page_str) - 1
            except Exception as e:
                print("Link çözümleme hatası:", e)

        path_check = Path(pdf_path)
        if not path_check.exists():
            reply = QMessageBox.question(
                self, t("msg_error"), t("err_file_relocate").format(pdf_path),
                QMessageBox.Yes | QMessageBox.No, QMessageBox.Yes
            )
            if reply == QMessageBox.Yes:
                new_filename, _ = QFileDialog.getOpenFileName(self, t("dlg_file_relocate"), str(Path.home()), "PDF Files (*.pdf)")
                if not new_filename: return
                self.db.update_pdf_path(pdf_path, new_filename)
                pdf_path = new_filename
            else:
                return

        self.current_pdf = pdf_path
        self.viewer.open_pdf(pdf_path)
        self.viewer.goto_position(page, x, y)
        self.update_page_label()
        self.setWindowTitle(f"{t('app_title')} - {Path(pdf_path).name}")
        self.right_stack.setCurrentIndex(1)
        self.load_pdf_bookmarks()

        self.note_edit.setReadOnly(True)
        html_note = current_note
        matches = re.findall(r'(\[\[([^\]]+)\]\])', current_note)
        for full_match, link_content in matches:
            html_link = f'<a href="{link_content}" style="color: #0078ff; font-weight: bold; text-decoration: none;">{full_match}</a>'
            html_note = html_note.replace(full_match, html_link)

        self.note_edit.setHtml(html_note.replace("\n", "<br>"))
        self.note_edit.setEnabled(True)
        self.btn_save_note.setEnabled(True)

    def enable_note_editing(self):
        if not self.active_bookmark_id or not self.note_edit.isReadOnly(): return
        bookmark_data = self.db.get_bookmark(self.active_bookmark_id)
        if bookmark_data and len(bookmark_data) > 7:
            current_note = bookmark_data[7] or ""
            self.note_edit.setReadOnly(False)
            self.note_edit.setTextInteractionFlags(Qt.TextEditable | Qt.TextSelectableByMouse)
            self.note_edit.setPlainText(current_note)

    def save_current_bookmark_note(self):
        if not self.active_bookmark_id: return
        try:
            notes_text = self.note_edit.toPlainText().strip()
            self.db.update_bookmark_notes(self.active_bookmark_id, notes_text)
            self.note_edit.setReadOnly(True)
            self.note_edit.setTextInteractionFlags(Qt.LinksAccessibleByMouse)

            html_note = notes_text
            matches = re.findall(r'(\[\[([^\]]+)\]\])', notes_text)
            for full_match, link_content in matches:
                html_link = f'<a href="{link_content}" style="color: #0078ff; font-weight: bold; text-decoration: none;">{full_match}</a>'
                html_note = html_note.replace(full_match, html_link)
            self.note_edit.setHtml(html_note.replace("\n", "<br>"))
        except Exception as e:
            QMessageBox.warning(self, t("msg_error"), t("err_save_note").format(e))

    def handle_knowledge_link(self, url):
        try:
            raw_str = url.toString().replace("link://", "").strip().rstrip("/")
            if not raw_str and url.isValid():
                raw_str = url.toDisplayString().replace("link://", "").strip().rstrip("/")
            if not raw_str: return

            if "#" not in raw_str:
                target_file, target_page = raw_str, 0
            else:
                target_file, page_str = raw_str.split("#", 1)
                target_file = target_file.strip()
                clean_page_str = "".join(filter(str.isdigit, page_str))
                target_page = int(clean_page_str) - 1 if clean_page_str else 0

            if self.current_pdf and (Path(self.current_pdf).name == target_file or self.current_pdf == target_file):
                self.viewer.goto_page(target_page)
                self.update_page_label()
                self.right_stack.setCurrentIndex(1)
                return

            full_pdf_path = None
            if self.current_pdf:
                sibling_path = Path(self.current_pdf).parent / target_file
                if sibling_path.exists() and sibling_path.is_file():
                    full_pdf_path = str(sibling_path)

            if not full_pdf_path:
                for item in getattr(self, "pdf_history", []):
                    p_str = item.get("path", "") if isinstance(item, dict) else item
                    if p_str and (p_str.endswith(target_file) or Path(p_str).name == target_file):
                        if Path(p_str).exists() and Path(p_str).is_file():
                            full_pdf_path = p_str
                            break

            if not full_pdf_path:
                QMessageBox.warning(self, t("msg_error"), t("err_target_pdf").format(target_file))
                return

            self.current_pdf = full_pdf_path
            self.viewer.open_pdf(full_pdf_path)
            self.viewer.goto_page(target_page)
            self.update_page_label()
            self.load_pdf_bookmarks()
            self.setWindowTitle(f"{t('app_title')} - {target_file}")
            self.right_stack.setCurrentIndex(1)
        except Exception as e:
            print("Knowledge link hatası:", e)

    def rename_group(self, item):
        if not item: return
        group_id = item.data(0, Qt.UserRole)
        if group_id is None: return
        name, ok = QInputDialog.getText(self, t("dlg_rename_group"), t("dlg_new_group_name"), text=item.text(0))
        if ok and name.strip():
            self.db.rename_group(group_id, name.strip())
            self.load_groups()

    def delete_group_recursive(self, item):
        if not item: return
        group_id = item.data(0, Qt.UserRole)
        if group_id is None: return
        reply = QMessageBox.question(self, t("dlg_del_group"), t("dlg_del_group_msg"))
        if reply == QMessageBox.Yes:
            self._delete_group_children(item)
            self.db.delete_group(group_id)
            self.load_groups()

    def _delete_group_children(self, item):
        for i in range(item.childCount()):
            child = item.child(i)
            gid = child.data(0, Qt.UserRole)
            if gid:
                self._delete_group_children(child)
                self.db.delete_group(gid)

    def add_subgroup(self, item):
        if not item: return
        group_id = item.data(0, Qt.UserRole)
        if group_id is None: return
        name, ok = QInputDialog.getText(self, t("dlg_new_subgroup"), t("dlg_subgroup_name"))
        if ok and name.strip():
            self.db.add_group(name.strip(), group_id)
            self.load_groups()

    def add_bookmark_from_selection(self, page, text, x, y):
        if self.selected_group_id is None:
            QMessageBox.warning(self, t("msg_warning"), t("warn_select_group"))
            return
        self.db.add_bookmark(text, self.current_pdf, page, self.selected_group_id, x, y)
        self.load_groups()
        self.update_page_label()

    def create_highlight(self, page_no, x0, y0, x1, y1):
        if not self.current_pdf or not self.viewer.doc: return
        try:
            page = self.viewer.doc[page_no]
            rect = pymupdf.Rect(x0, y0, x1, y1)
            page.add_highlight_annot(rect)
            self.viewer.doc.saveIncr()
            self.viewer.render_page()
        except Exception as e:
            QMessageBox.warning(self, t("msg_error"), t("err_highlight_readonly").format(e))

    def delete_highlight(self, page_no, x, y):
        if not self.current_pdf or not self.viewer.doc: return
        try:
            page = self.viewer.doc[page_no]
            point = pymupdf.Point(x, y)
            annot = page.first_annot
            while annot:
                next_annot = annot.next
                if annot.rect.contains(point):
                    page.delete_annot(annot)
                    self.viewer.doc.saveIncr()
                    self.viewer.render_page()
                    return
                annot = next_annot
        except Exception as e:
            QMessageBox.warning(self, t("msg_error"), t("err_highlight_readonly").format(e))

    def print_pdf(self):
        if not self.current_pdf: return
        items = [t("print_opt1"), t("print_opt2"), t("print_opt3")]
        choice, ok = QInputDialog.getItem(self, t("print_title"), t("print_q1"), items, 0, False)
        if not ok: return
        if choice == t("print_opt3"): self.print_all_pages()
        elif choice == t("print_opt1"): self.print_current_page()
        elif choice == t("print_opt2"): self.print_page_range()

    def _get_default_printer(self):
        try:
            result = subprocess.run(["lpstat", "-d"], capture_output=True, text=True, check=True)
            output = result.stdout.strip()
            if ":" in output: return output.split(":", 1)[1].strip()
        except Exception: pass
        return None

    def print_all_pages(self):
        if not shutil.which("lp"):
            QMessageBox.warning(self, t("print_err"), t("print_err_missing_tool").format("lp"))
            return
        try:
            printer = self._get_default_printer()
            cmd = ["lp"]
            if printer: cmd.extend(["-d", printer])
            cmd.append(self.current_pdf)
            subprocess.run(cmd, check=True)
            info = f"({printer})" if printer else t("print_default_printer")
            QMessageBox.information(self, t("print_title"), t("print_all_success").format(info))
        except Exception as e:
            QMessageBox.warning(self, t("print_err"), str(e))

    def print_page_range(self):
        if not shutil.which("qpdf"):
            QMessageBox.warning(self, t("print_err"), t("print_err_missing_tool").format("qpdf"))
            return
        if not shutil.which("lp"):
            QMessageBox.warning(self, t("print_err"), t("print_err_missing_tool").format("lp"))
            return

        pages, ok = QInputDialog.getText(self, t("print_title"), t("print_range_q"))
        if not ok or not pages.strip(): return
        try:
            with tempfile.TemporaryDirectory() as tmpdir:
                output_pdf = Path(tmpdir) / "range_output.pdf"
                subprocess.run(["qpdf", self.current_pdf, "--pages", self.current_pdf, pages, "--", str(output_pdf)], check=True)
                printer = self._get_default_printer()
                cmd = ["lp"]
                if printer: cmd.extend(["-d", printer])
                cmd.append(str(output_pdf))
                subprocess.run(cmd, check=True)
            info = f"({printer})" if printer else t("print_default_printer")
            QMessageBox.information(self, t("print_title"), t("print_range_success").format(pages, info))
        except Exception as e:
            QMessageBox.warning(self, t("print_err"), t("print_range_err").format(e))

    def print_current_page(self):
        if not shutil.which("pdfseparate"):
            QMessageBox.warning(self, t("print_err"), t("print_err_missing_tool").format("pdfseparate (poppler)"))
            return
        if not shutil.which("lp"):
            QMessageBox.warning(self, t("print_err"), t("print_err_missing_tool").format("lp"))
            return

        page = self.viewer.current_page + 1
        try:
            with tempfile.TemporaryDirectory() as tmpdir:
                actual_file = Path(tmpdir) / f"page_{page}.pdf"
                output_template = Path(tmpdir) / "page_%d.pdf"
                subprocess.run(["pdfseparate", "-f", str(page), "-l", str(page), self.current_pdf, str(output_template)], check=True)
                if actual_file.exists():
                    printer = self._get_default_printer()
                    cmd = ["lp"]
                    if printer: cmd.extend(["-d", printer])
                    cmd.append(str(actual_file))
                    subprocess.run(cmd, check=True)
                    info = f"({printer})" if printer else t("print_default_printer")
                    QMessageBox.information(self, t("print_title"), t("print_page_success").format(page, info))
                else:
                    raise FileNotFoundError(t("print_page_err_file"))
        except Exception as e:
            QMessageBox.warning(self, t("print_err"), t("print_page_err").format(e))

    def goto_page(self):
        if not self.viewer.doc: return
        try:
            page_no = int(self.page_edit.text())
            if 1 <= page_no <= self.viewer.doc.page_count:
                self.viewer.current_page = page_no - 1
                self.viewer.render_page()
        except ValueError: pass
        finally: self.update_page_label()

    def export_bookmarks_to_markdown(self):
        if not self.current_pdf:
            QMessageBox.warning(self, t("msg_warning"), t("export_warn_open"))
            return
        file_name = Path(self.current_pdf).stem
        default_save = PROJE_ANA_DIZINI / f"{file_name}_Calisma_Notlari.md"
        save_filename, _ = QFileDialog.getSaveFileName(self, t("export_save_title"), str(default_save), "Markdown Files (*.md)")
        if not save_filename: return

        try:
            bookmarks = self.db.get_bookmarks()
            current_bookmarks = [b for b in bookmarks if b[2] == self.current_pdf]
            if not current_bookmarks:
                QMessageBox.information(self, t("msg_info"), t("export_no_bookmarks"))
                return
            groups = {g[0]: g[1] for g in self.db.get_groups()}
            current_bookmarks.sort(key=lambda x: x[3])

            md = t("export_md_title")
            md += t("export_md_doc").format(Path(self.current_pdf).name)
            md += t("export_md_path").format(self.current_pdf)
            md += t("export_md_date").format(datetime.datetime.now().strftime('%Y-%m-%d %H:%M'))
            md += t("export_md_header")

            for bookmark in current_bookmarks:
                title, page, group_id = bookmark[1], bookmark[3], bookmark[4]
                group_name = groups.get(group_id, t("export_general"))
                md += f"### 📘 {title}\n"
                md += t("export_md_group").format(group_name, page + 1)
                note_text = bookmark[7].strip() if len(bookmark) > 7 and bookmark[7] else None
                if note_text:
                    md += t("export_md_note_title")
                    indented_note = "\n".join([f"  > {line}" for line in note_text.split("\n")])
                    md += f"{indented_note}\n"
                else:
                    md += t("export_md_no_note")
                md += "\n"

            with open(save_filename, "w", encoding="utf-8") as f:
                f.write(md)
            QMessageBox.information(self, t("msg_success"), t("export_success").format(Path(save_filename).name))
        except Exception as e:
            QMessageBox.warning(self, t("msg_error"), t("export_err").format(e))

    def export_table_to_excel(self, page_no, x0, y0, x1, y1):
        if not self.current_pdf or not self.viewer.doc: return
        try:
            import pandas as pd
        except ImportError:
            QMessageBox.warning(self, t("exc_lib_err"), t("exc_lib_msg"))
            return
        try:
            page = self.viewer.doc[page_no]
            rect = pymupdf.Rect(x0, y0, x1, y1)
            tables = page.find_tables(clip=rect)

            if not tables or not tables.tables:
                QMessageBox.warning(self, t("exc_notable_title"), t("exc_notable_msg"))
                return
            data = tables.tables[0].extract()
            if not data:
                QMessageBox.warning(self, t("exc_empty_title"), t("exc_empty_msg"))
                return

            df = pd.DataFrame(data).fillna('').astype(str).replace(['NaN', 'None'], '', regex=True)
            cleaned_rows = df.apply(lambda row: [str(val).replace('\n', ' ').strip() for val in row if str(val).strip() != ''], axis=1).tolist()
            if not cleaned_rows: return
            max_cols = max(len(row) for row in cleaned_rows if row)
            final_data = []
            for i, row in enumerate(cleaned_rows):
                if not row: final_data.append([""] * max_cols); continue
                if len(row) < max_cols:
                    if i == 0: row = [""] * (max_cols - len(row)) + row
                    else: row = row + [""] * (max_cols - len(row))
                final_data.append(row)

            df_clean = pd.DataFrame(final_data)
            docs_dir = Path.home() / "Documents"
            if not docs_dir.exists(): docs_dir = Path.home()
            default_save = str(docs_dir / f"{Path(self.current_pdf).stem}_Tablo_Sayfa{page_no + 1}.xlsx")
            save_filename, _ = QFileDialog.getSaveFileName(self, t("exc_save_title"), default_save, "Excel Files (*.xlsx)")

            if save_filename:
                df_clean.to_excel(save_filename, index=False, header=False)
                try:
                    from openpyxl import load_workbook
                    from openpyxl.styles import Font
                    wb = load_workbook(save_filename)
                    ws = wb.active
                    serif_font = Font(name="Times New Roman", size=11)
                    for row in ws.iter_rows():
                        for cell in row: cell.font = serif_font
                    wb.save(save_filename)
                except Exception: pass

                msg_box = QMessageBox(self)
                msg_box.setWindowTitle(t("msg_success"))
                msg_box.setText(t("exc_success_msg").format(Path(save_filename).name))
                msg_box.setIcon(QMessageBox.Information)
                btn_open = msg_box.addButton(t("exc_btn_open"), QMessageBox.ActionRole)
                msg_box.addButton(t("exc_btn_ok"), QMessageBox.AcceptRole)
                msg_box.exec()
                if msg_box.clickedButton() == btn_open:
                    QDesktopServices.openUrl(QUrl.fromLocalFile(save_filename))
        except Exception as e:
            QMessageBox.warning(self, t("msg_error"), t("exc_err").format(e))

    def add_cross_pdf_link(self, item):
        bid = item.data(0, Qt.UserRole + 1)
        if not bid: return
        pdf_options = [t("link_wiz_newfile")]
        history = getattr(self, "pdf_history", [])
        pdf_options.extend([Path(d.get("path")).name for d in history if isinstance(d, dict) and d.get("path")])
        target_file, ok1 = QInputDialog.getItem(self, t("link_wiz_title"), t("link_wiz_q1"), pdf_options, 0, False)
        if not ok1 or not target_file: return
        final_file_name = target_file
        if target_file == t("link_wiz_newfile"):
            filename, _ = QFileDialog.getOpenFileName(self, t("link_wiz_open"), str(Path.home()), "PDF Files (*.pdf)")
            if not filename: return
            final_file_name = Path(filename).name
            self.pdf_history = [i for i in self.pdf_history if (i.get("path") if isinstance(i, dict) else i) != filename]
            self.pdf_history.insert(0, {"path": filename, "page": 0, "date": datetime.datetime.now().strftime("%Y-%m-%d %H:%M")})

        target_page, ok2 = QInputDialog.getInt(self, t("link_wiz_title"), t("link_wiz_q2").format(final_file_name), 1, 1, 9999, 1)
        if not ok2: return
        try:
            self.note_edit.setReadOnly(False)
            updated_text = self.note_edit.toPlainText() + f" [[{final_file_name}#{target_page}]]"
            self.db.update_bookmark_notes(bid, updated_text)
            self.note_edit.setReadOnly(True)
            html_note = updated_text
            for full_match, link_content in re.findall(r'(\[\[([^\]]+)\]\])', updated_text):
                html_link = f'<a href="{link_content}" style="color: #0078ff; font-weight: bold; text-decoration: none;">{full_match}</a>'
                html_note = html_note.replace(full_match, html_link)
            self.note_edit.setHtml(html_note.replace("\n", "<br>"))
        except Exception as e:
            QMessageBox.warning(self, t("msg_error"), str(e))

if __name__ == "__main__":
    app = QApplication(sys.argv)
    app.setApplicationName("PDF_Network")
    app.setDesktopFileName("PDF_Network")
    window = MainWindow()
    if len(sys.argv) > 1:
        initial_pdf = sys.argv[1]
        if Path(initial_pdf).exists():
            window.current_pdf = initial_pdf
            window.viewer.open_pdf(initial_pdf)
            window.update_page_label()
            window.load_pdf_bookmarks()
            window.setWindowTitle(f"{t('app_title')} - {Path(initial_pdf).name}")
            window.right_stack.setCurrentIndex(1)
    window.show()
    sys.exit(app.exec())
